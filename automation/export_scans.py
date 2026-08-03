"""Export everything the engine has scanned to a single .xlsx.

For the interim setup where the ERP cannot reach the engine (no tunnel yet), so
finance gets a file to look at instead of a web page. Reads the SQLite audit
trail directly - the engine does not need to be running.

    python export_scans.py                     # -> data/scanned-bills-<date>.xlsx
    python export_scans.py --out C:\\path\\x.xlsx
    python export_scans.py --status review     # only bills awaiting a human

Sheet 1  one row per scanned page, with the fields OCR extracted
Sheet 2  the full OCR text per page
Sheet 3  exports already generated for Tally, and where the XML sits
"""
from __future__ import annotations

import argparse
import datetime as _dt
import os
import re
import sqlite3
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
except ImportError:
    sys.exit("openpyxl is missing - run: python -m pip install -r requirements.txt")

BASE = os.path.dirname(os.path.abspath(__file__))
DB = os.path.join(BASE, "data", "finance.db")

# openpyxl refuses these outright (IllegalCharacterError), and Tesseract does emit
# them - \x0c in particular, on multi-column receipts. app/tally.py strips the same
# ranges before building XML; without it one bad byte kills the whole export.
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def clean(v):
    """Strip control characters openpyxl will not accept. Non-strings pass through."""
    return _ILLEGAL.sub("", v) if isinstance(v, str) else v

FIELDS = [
    ("vendor_name", "Vendor"), ("vendor_gstin", "Vendor GSTIN"),
    ("invoice_no", "Invoice no"), ("invoice_date", "Invoice date"),
    ("taxable_value", "Taxable"), ("cgst", "CGST"), ("sgst", "SGST"),
    ("igst", "IGST"), ("round_off", "Round off"), ("net_amount", "Net amount"),
    ("hsn_sac", "HSN/SAC"), ("ledger", "Ledger"), ("narration", "Narration"),
    ("confidence", "Match conf"), ("arithmetic_ok", "Arithmetic OK"),
]


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", help="output .xlsx path")
    ap.add_argument("--status", help="only bills with this status (e.g. review, posted)")
    args = ap.parse_args()

    if not os.path.exists(DB):
        sys.exit(f"no database at {DB} - nothing has been scanned yet")

    out = args.out or os.path.join(BASE, "data", f"scanned-bills-{_dt.date.today():%Y-%m-%d}.xlsx")
    conn = sqlite3.connect(DB)
    conn.row_factory = sqlite3.Row

    have = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    ecols = {d[1] for d in conn.execute("PRAGMA table_info(extractions)")} if "extractions" in have else set()
    picked = [(c, label) for c, label in FIELDS if c in ecols]

    sel = "SELECT b.*, " + ", ".join(f"e.{c} AS x_{c}" for c, _ in picked) if picked else "SELECT b.*"
    sql = f"{sel} FROM bills b LEFT JOIN extractions e ON e.bill_id = b.id"
    params: list = []
    if args.status:
        sql += " WHERE b.status = ?"
        params.append(args.status)
    sql += " ORDER BY b.id"
    rows = conn.execute(sql, params).fetchall()
    if not rows:
        sys.exit("no bills matched - nothing to export")

    wb = Workbook()

    ws = wb.active
    ws.title = "Scanned bills"
    head = ["Bill ID", "Source file", "Page", "Person", "Status", "Quality",
            "OCR conf %", "OCR engine", "Scanned at"] + [label for _, label in picked]
    ws.append(head)
    for r in rows:
        conf = r["ocr_confidence"]
        ws.append([clean(v) for v in ([
            r["id"], os.path.basename(r["source_file"] or ""), r["page_no"], r["person"],
            r["status"], r["quality_verdict"],
            round(conf, 1) if conf is not None else None,
            f'{r["ocr_engine"] or ""}/{r["ocr_variant"] or ""}'.strip("/"),
            r["created_at"],
        ] + [r[f"x_{c}"] for c, _ in picked])])

    w2 = wb.create_sheet("OCR text")
    w2.append(["Bill ID", "Source file", "Page", "OCR text"])
    for r in rows:
        w2.append([r["id"], clean(os.path.basename(r["source_file"] or "")),
                   r["page_no"], clean(r["ocr_text"] or "")])
    w2.column_dimensions["D"].width = 120
    for row in w2.iter_rows(min_row=2, min_col=4, max_col=4):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    w3 = wb.create_sheet("Tally exports")
    if "exports" in have:
        ecol = [d[1] for d in conn.execute("PRAGMA table_info(exports)")]
        w3.append(ecol + ["XML on disk?"])
        xdir = os.path.join(BASE, "data", "tally_export")
        for e in conn.execute("SELECT * FROM exports ORDER BY id"):
            # api.py writes the batch to <ref>.xml, and ref carries no extension -
            # matching on a value ending in .xml never hits.
            ref = e["ref"] if "ref" in ecol else None
            on_disk = bool(ref) and os.path.exists(os.path.join(xdir, f"{ref}.xml"))
            # the xml column holds the whole document; a cell tops out at 32767 chars
            vals = [("<%d chars>" % len(e[c])) if c == "xml" and isinstance(e[c], str)
                    else clean(e[c]) for c in ecol]
            w3.append(vals + ["yes" if on_disk else "no"])
    else:
        w3.append(["no exports table"])

    for sheet in (ws, w2, w3):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
    for col, width in zip("ABCDEFGHI", [8, 30, 6, 9, 9, 9, 11, 16, 20]):
        ws.column_dimensions[col].width = width
    for col, width in zip("ABC", [8, 30, 6]):
        w2.column_dimensions[col].width = width

    wb.save(out)
    by_status: dict[str, int] = {}
    for r in rows:
        by_status[r["status"]] = by_status.get(r["status"], 0) + 1
    print(f"wrote {out}")
    print(f"  {len(rows)} scanned pages  {by_status}")
    print("  note: only confirmed bills can be exported to Tally XML")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
